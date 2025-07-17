from openpyxl import load_workbook

def is_row_hidden(sheet, row_index):
    return sheet.row_dimensions[row_index].hidden if row_index in sheet.row_dimensions else False

def get_merged_cells(sheet):
    return {
        (r, c)
        for merged_range in sheet.merged_cells.ranges
        for r in range(merged_range.min_row, merged_range.max_row + 1)
        for c in range(merged_range.min_col, merged_range.max_col + 1)
    }
    #merged_set = set()
    #for merged_range in sheet.merged_cells.ranges:
    #    min_col, min_row, max_col, max_row = merged_range.bounds
    #    for r in range(min_row, max_row + 1):
    #        for c in range(min_col, max_col + 1):
    #            merged_set.add((r, c))
    #return merged_set

def header_finder(sheet):
    for row in sheet.iter_rows():
        for cell in row:
            if str(cell.value).lower() == "артикул":
                print(f"Заголовок Артикул найден в строке {cell.row} и столбце {cell.column}")
                return cell.row, cell.column
    print("Артикулы не найдены")
    return None

def find_articles(file_path):
    articles = list()

    wb = load_workbook(filename=file_path)
    ws = wb.active

    header = header_finder(ws)
    # merged_cells = get_merged_cells(ws)


    if not header:
        return
    header_row, header_col = header
    for row in ws.iter_rows(min_row=header_row + 1, min_col=header_col, max_col=header_col):
        cell = row[0]
        if is_row_hidden(ws, cell.row) or cell.value.split()[0] == "Склад":
            continue
        articles.append(cell.value)
    return articles


# pyside